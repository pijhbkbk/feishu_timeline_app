from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from collections import Counter
import datetime as dt
import hashlib, json, os, posixpath
import openpyxl
from openpyxl.utils.cell import get_column_letter, coordinate_to_tuple, range_boundaries

from io import BytesIO
from source_guard import SourceLimits, SourceValidationError, validate_workbook_source
from io_utils import write_json

NS = {'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
def typed(v):
    if isinstance(v,(dt.datetime,dt.date,dt.time)): return {'type':type(v).__name__,'value':v.isoformat()}
    return {'type':type(v).__name__,'value':v}

def extract_workbook(source: Path, output: Path, limits: SourceLimits=SourceLimits(), writer=write_json):
    SOURCE=source
    OUT=output
    guard=validate_workbook_source(SOURCE, limits)
    with SOURCE.open('rb') as f: raw=f.read(limits.max_file_bytes+1)
    if len(raw)>limits.max_file_bytes or hashlib.sha256(raw).hexdigest()!=guard['sourceSHA']: raise SourceValidationError('Source changed during validation')
    before=guard['sourceSHA']
    wb=openpyxl.load_workbook(BytesIO(raw),data_only=False,read_only=False,keep_links=False)
    cached=openpyxl.load_workbook(BytesIO(raw),data_only=True,read_only=False,keep_links=False)
    book={'sourcePath':str(SOURCE),'sourceSha256':before,'sourceBytes':len(raw),'epoch':wb.epoch.isoformat(),
          'sheetNames':wb.sheetnames,'definedNames':[dict(x) for x in wb.defined_names.values()], 'sheets':[]}
    summary=[]; row_dump=[]
    with ZipFile(BytesIO(raw)) as z:
        rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        relmap={e.attrib['Id']:posixpath.normpath(posixpath.join('xl',e.attrib['Target'].lstrip('/'))) if not e.attrib['Target'].startswith('/') else e.attrib['Target'].lstrip('/') for e in rels}
        xbook=ET.fromstring(z.read('xl/workbook.xml'))
        sheetparts={e.attrib['name']:relmap[e.attrib[REL]] for e in xbook.find('x:sheets',NS)}
        rawdir=OUT/'source-ooxml'; rawdir.mkdir(exist_ok=True,mode=0o700)
        book['archiveParts']=[]
        for name in z.namelist():
            part=z.read(name)
            is_dir=z.getinfo(name).is_dir()
            book['archiveParts'].append({'name':name,'bytes':len(part),'sha256':hashlib.sha256(part).hexdigest(),'isDirectory':is_dir})
            p=rawdir/name
            if is_dir:
                p.mkdir(parents=True,exist_ok=True,mode=0o700)
            else:
                p.parent.mkdir(parents=True,exist_ok=True,mode=0o700); p.write_bytes(part);p.chmod(0o600)
        for ws in wb.worksheets:
            xml=ET.fromstring(z.read(sheetparts[ws.title])); native={e.attrib['r']:e for e in xml.findall('.//x:sheetData/x:row/x:c',NS)}
            # openpyxl's populated-cell extent can be smaller than the declared source
            # rectangle. Preserve declared blank rows/columns instead of trimming them.
            dimension=xml.find('x:dimension',NS)
            declared=range_boundaries(dimension.attrib['ref']) if dimension is not None else (1,1,1,1)
            rect_row=max(ws.max_row,declared[3],max(ws.row_dimensions.keys(),default=1))
            rect_col=max(ws.max_column,declared[2])
            ws.cell(row=rect_row,column=rect_col)
            merged={}
            for ran in ws.merged_cells.ranges:
                for row in ws.iter_rows(min_row=ran.min_row,max_row=ran.max_row,min_col=ran.min_col,max_col=ran.max_col):
                    for c in row: merged[c.coordinate]={'range':str(ran),'anchor':f'{get_column_letter(ran.min_col)}{ran.min_row}'}
            sh={'name':ws.title,'state':ws.sheet_state,'dimension':ws.calculate_dimension(),'maxRow':ws.max_row,'maxColumn':ws.max_column,
                'xmlPart':sheetparts[ws.title], 'xmlDimension':xml.find('x:dimension',NS).attrib if xml.find('x:dimension',NS) is not None else None,
                'mergedRanges':[str(x) for x in ws.merged_cells.ranges],
                'rowDimensions':{str(k):dict(v) for k,v in ws.row_dimensions.items()},
                'columnDimensions':{str(k):dict(v) for k,v in ws.column_dimensions.items()},
                'freezePanes':str(ws.freeze_panes) if ws.freeze_panes else None,
                'autoFilter':str(ws.auto_filter.ref) if ws.auto_filter else None,
                'tables':list(ws.tables),'dataValidations':[str(x) for x in ws.data_validations.dataValidation],
                'printArea':str(ws.print_area),'rows':[]}
            kinds=Counter(); nonemptyrows=0; form=0; zeros=0; dates=0; datedvalues=0; comments=0; styles=Counter(); labels=[]
            for cells in ws.iter_rows():
                rowdata=[]; printed=[]
                for c in cells:
                    e=native.get(c.coordinate); v=e.find('x:v',NS) if e is not None else None; f=e.find('x:f',NS) if e is not None else None
                    is_date=bool(getattr(c,'is_date',False)); comment=getattr(c,'comment',None); hyperlink=getattr(c,'hyperlink',None)
                    entry={'coordinate':c.coordinate,'row':c.row,'column':c.column,'value':typed(c.value),'dataType':c.data_type,
                           'numberFormat':c.number_format,'isDate':is_date,'styleId':c.style_id,
                           'xmlPresent':e is not None,'xmlAttributes':e.attrib if e is not None else None,
                           'xmlValuePresent':v is not None,'xmlValue':v.text if v is not None else None,
                           'formula':{'text':f.text,'attributes':f.attrib} if f is not None else None,
                           'cachedValue':typed(cached[ws.title][c.coordinate].value) if f is not None else None,
                           'merge':merged.get(c.coordinate),
                           'comment':{'text':comment.text,'author':comment.author} if comment else None,
                           'hyperlink':hyperlink.target if hyperlink else None}
                    if c.value is None:
                        entry['blankKind']='merged-child' if c.coordinate in merged and merged[c.coordinate]['anchor']!=c.coordinate else ('explicit-blank' if e is not None else 'implicit-blank')
                    elif isinstance(c.value,str) and c.value=='': entry['blankKind']='empty-string'
                    elif isinstance(c.value,str) and not c.value.strip(): entry['blankKind']='whitespace-string'
                    else: entry['blankKind']=None
                    rowdata.append(entry); kinds[entry['blankKind'] or 'value']+=1; styles[c.number_format]+=1
                    if c.value is not None:
                        printed.append(f'{c.coordinate}={json.dumps(typed(c.value)["value"],ensure_ascii=False)}'+(f' [format:{c.number_format}]' if is_date else ''))
                    form+=f is not None; comments+=comment is not None; dates+=is_date
                    datedvalues+=isinstance(c.value,(dt.date,dt.datetime,dt.time))
                    zeros+=isinstance(c.value,(int,float)) and not isinstance(c.value,bool) and c.value==0
                sh['rows'].append({'row':cells[0].row,'hidden':bool(ws.row_dimensions[cells[0].row].hidden),'cells':rowdata})
                if printed: nonemptyrows+=1
                row_dump.append(f'[{ws.title}] row {cells[0].row}: '+(' | '.join(printed) if printed else '<ENTIRE ROW BLANK>'))
            sh['counts']={'rectangularCells':ws.max_row*ws.max_column,'xmlCells':len(native),'nonemptyRows':nonemptyrows,'blankKinds':dict(kinds),'formulas':form,'dateFormattedCells':dates,'dateValues':datedvalues,'zeroCells':zeros,'comments':comments,'numberFormats':dict(styles)}
            book['sheets'].append(sh)
            summary.append({k:sh[k] for k in ['name','state','dimension','maxRow','maxColumn','mergedRanges','freezePanes','autoFilter','tables','counts']})
    with SOURCE.open('rb') as f: after=f.read(limits.max_file_bytes+1)
    book['sourceAfterSha256']=hashlib.sha256(after).hexdigest()
    assert before==book['sourceAfterSha256'],'Source unexpectedly changed'
    book['guardSummary']=guard
    writer(OUT/'workbook-extraction.json',book)
    writer(OUT/'safe-structure-counts.json',{'sourceSha256':before,'sourceUnchanged':True,'sheets':summary})
    (OUT/'all-rows.txt').write_text('\n'.join(row_dump),encoding='utf-8');(OUT/'all-rows.txt').chmod(0o600)
    return book
